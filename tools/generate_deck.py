"""
Generate a professional, modern 16:9 widescreen PowerPoint presentation for Parity.
Uses dark aesthetic with luminous accents, metric cards, workflow diagrams, and structured feature blocks.
"""
import sys
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# ── Color Palette ───────────────────────────────────────────────────────────
BG_DARK = RGBColor(11, 15, 25)         # #0B0F19 - Deep dark canvas
CARD_BG = RGBColor(24, 34, 53)         # #182235 - Surface card background
CARD_BORDER = RGBColor(51, 65, 85)     # #334155 - Subtle border
TEXT_WHITE = RGBColor(248, 250, 252)   # #F8FAFC - Primary text
TEXT_MUTED = RGBColor(148, 163, 184)   # #94A3B8 - Secondary text
TEXT_DIM = RGBColor(100, 116, 139)     # #64748B - Caption / dim text

ACCENT_BLUE = RGBColor(56, 189, 248)   # #38BDF8 - Sky Cyan
ACCENT_PURPLE = RGBColor(168, 85, 247) # #A855F7 - Purple
ACCENT_GREEN = RGBColor(52, 211, 153)  # #34D399 - Emerald
ACCENT_AMBER = RGBColor(251, 191, 36)  # #FBBF24 - Amber
ACCENT_ROSE = RGBColor(244, 63, 94)    # #F43F5E - Rose / Red

FONT_HEADING = "Segoe UI"
FONT_BODY = "Segoe UI"
FONT_CODE = "Consolas"

def create_presentation() -> Presentation:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6] # Blank slide

    def set_slide_background(slide):
        bg = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height
        )
        bg.fill.solid()
        bg.fill.fore_color.rgb = BG_DARK
        bg.line.fill.background()
        return bg

    def add_header(slide, title_text: str, category_text: str = "PARITY ENVIRONMENT DIAGNOSTICS"):
        # Category Tag
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.45), Inches(11.5), Inches(0.35))
        tf_cat = cat_box.text_frame
        tf_cat.word_wrap = True
        tf_cat.margin_left = tf_cat.margin_right = tf_cat.margin_top = tf_cat.margin_bottom = 0
        p_cat = tf_cat.paragraphs[0]
        p_cat.text = category_text.upper()
        p_cat.font.name = FONT_HEADING
        p_cat.font.size = Pt(10.5)
        p_cat.font.bold = True
        p_cat.font.color.rgb = ACCENT_BLUE

        # Main Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.75), Inches(11.5), Inches(0.7))
        tf = title_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.text = title_text
        p.font.name = FONT_HEADING
        p.font.size = Pt(26)
        p.font.bold = True
        p.font.color.rgb = TEXT_WHITE

    def add_card(slide, left, top, width, height, bg_color=CARD_BG, border_color=CARD_BORDER):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = bg_color
        if border_color:
            card.line.color.rgb = border_color
            card.line.width = Pt(1)
        else:
            card.line.fill.background()
        return card

    # ═════════════════════════════════════════════════════════════════════════
    # SLIDE 1: Title Slide
    # ═════════════════════════════════════════════════════════════════════════
    s1 = prs.slides.add_slide(blank_layout)
    set_slide_background(s1)

    # Accent decorative glow box
    glow = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.8), Inches(0.08), Inches(3.8))
    glow.fill.solid()
    glow.fill.fore_color.rgb = ACCENT_BLUE
    glow.line.fill.background()

    # Title content frame
    t_box = s1.shapes.add_textbox(Inches(1.1), Inches(1.7), Inches(11.2), Inches(4.0))
    tf1 = t_box.text_frame
    tf1.word_wrap = True

    p0 = tf1.paragraphs[0]
    p0.text = "DEVELOPER TOOLING & RUNTIME DIAGNOSTICS"
    p0.font.name = FONT_HEADING
    p0.font.size = Pt(12)
    p0.font.bold = True
    p0.font.color.rgb = ACCENT_BLUE
    p0.space_after = Pt(14)

    p1 = tf1.add_paragraph()
    p1.text = "PARITY"
    p1.font.name = FONT_HEADING
    p1.font.size = Pt(56)
    p1.font.bold = True
    p1.font.color.rgb = TEXT_WHITE
    p1.space_after = Pt(10)

    p2 = tf1.add_paragraph()
    p2.text = "“Make Every Machine Match.”"
    p2.font.name = FONT_HEADING
    p2.font.size = Pt(22)
    p2.font.italic = True
    p2.font.color.rgb = ACCENT_PURPLE
    p2.space_after = Pt(18)

    p3 = tf1.add_paragraph()
    p3.text = "A deterministic CLI tool and interactive web engine that diagnoses why Python projects fail on one machine while succeeding on another, explains the exact root cause, and applies verified fixes safely."
    p3.font.name = FONT_BODY
    p3.font.size = Pt(14.5)
    p3.font.color.rgb = TEXT_MUTED

    # Bottom pill badges
    badges = [
        ("Python 3.11+", ACCENT_BLUE),
        ("AST Static Analysis", ACCENT_PURPLE),
        ("Zero Execution Security", ACCENT_GREEN),
        ("Web GUI + CLI", ACCENT_AMBER),
        ("46 Verified Tests", ACCENT_ROSE),
    ]
    bx = Inches(1.1)
    by = Inches(5.9)
    for text, color in badges:
        card = add_card(s1, bx, by, Inches(2.2), Inches(0.5), bg_color=CARD_BG, border_color=color)
        t = s1.shapes.add_textbox(bx, by, Inches(2.2), Inches(0.5))
        tf = t.text_frame
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        p.text = text
        p.font.name = FONT_HEADING
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = color
        bx += Inches(2.35)

    # ═════════════════════════════════════════════════════════════════════════
    # SLIDE 2: The Problem ("Works on My Machine")
    # ═════════════════════════════════════════════════════════════════════════
    s2 = prs.slides.add_slide(blank_layout)
    set_slide_background(s2)
    add_header(s2, "The Real-World Dilemma: Silent Environment Drift", "THE CHALLENGE")

    cards_data = [
        ("Undeclared Imports", "Developers import modules present in their local site-packages but omit them from requirements.txt or pyproject.toml. The project runs locally but instantly fails in CI or Docker.", ACCENT_ROSE),
        ("Version Drift & Disjoint Pins", "Loose specifiers (e.g. >=2.0) pull breaking upstream versions on other machines. Multiple files or subprojects declare contradictory pins (==1.4 vs ==2.0).", ACCENT_AMBER),
        ("Platform & Marker Gaps", "PEP 508 markers (sys_platform == 'win32') or architecture-specific wheels cause silent runtime missing-module exceptions when code crosses operating systems.", ACCENT_BLUE),
        ("Missing Native Tools", "Python packages often wrap C/C++ or system tools. Missing compilers (gcc, MSVC), build systems (cmake, make), or binaries (git, ffmpeg) break builds cryptically.", ACCENT_PURPLE),
    ]

    for i, (title, desc, color) in enumerate(cards_data):
        cx = Inches(0.8 + (i % 2) * 5.95)
        cy = Inches(1.7 + (i // 2) * 2.5)
        add_card(s2, cx, cy, Inches(5.7), Inches(2.2))

        # Accent tab
        tab = s2.shapes.add_shape(MSO_SHAPE.RECTANGLE, cx + Inches(0.3), cy + Inches(0.3), Inches(0.08), Inches(0.5))
        tab.fill.solid()
        tab.fill.fore_color.rgb = color
        tab.line.fill.background()

        tb = s2.shapes.add_textbox(cx + Inches(0.55), cy + Inches(0.25), Inches(4.8), Inches(1.7))
        tf = tb.text_frame
        tf.word_wrap = True
        p_t = tf.paragraphs[0]
        p_t.text = title
        p_t.font.name = FONT_HEADING
        p_t.font.size = Pt(16)
        p_t.font.bold = True
        p_t.font.color.rgb = TEXT_WHITE
        p_t.space_after = Pt(8)

        p_d = tf.add_paragraph()
        p_d.text = desc
        p_d.font.name = FONT_BODY
        p_d.font.size = Pt(12)
        p_d.font.color.rgb = TEXT_MUTED

    # ═════════════════════════════════════════════════════════════════════════
    # SLIDE 3: The Parity Solution
    # ═════════════════════════════════════════════════════════════════════════
    s3 = prs.slides.add_slide(blank_layout)
    set_slide_background(s3)
    add_header(s3, "Parity's Core Pillars & Value Proposition", "SOLUTION OVERVIEW")

    pillars = [
        ("AST Static Analysis", "Zero Untrusted Execution", "Scans Python code trees via standard AST parsing. Categorizes imports into Runtime, Dev/Test, and Optional without running arbitrary setup scripts or untrusted code.", ACCENT_BLUE),
        ("Deep Fingerprinting", "Deterministic System Snapshots", "Extracts active Python runtime, architecture, platform tags, complete package lists, and native compiler status into shareable, redacted JSON snapshots.", ACCENT_PURPLE),
        ("Rule-Based Diagnosis", "Actionable Confidence Scoring", "Evaluates formal diagnostic rules (R1-R7), pinpointing exact files, line numbers, severity ratings (High/Med/Low), and confidence metrics for every discrepancy.", ACCENT_GREEN),
        ("Automated Remediation", "CLI & Interactive Dashboard", "Builds safe, tiered remediation plans—appending missing packages, running pip commands, installing native tools via winget, and updating system PATH with 1 click.", ACCENT_AMBER),
    ]

    col_w = Inches(2.78)
    for i, (title, sub, body, color) in enumerate(pillars):
        px = Inches(0.8 + i * 2.95)
        py = Inches(1.7)
        add_card(s3, px, py, col_w, Inches(5.1))

        # Top border accent
        top_bar = s3.shapes.add_shape(MSO_SHAPE.RECTANGLE, px, py, col_w, Inches(0.08))
        top_bar.fill.solid()
        top_bar.fill.fore_color.rgb = color
        top_bar.line.fill.background()

        tb = s3.shapes.add_textbox(px + Inches(0.2), py + Inches(0.3), col_w - Inches(0.4), Inches(4.5))
        tf = tb.text_frame
        tf.word_wrap = True

        p1 = tf.paragraphs[0]
        p1.text = f"0{i+1}"
        p1.font.name = FONT_HEADING
        p1.font.size = Pt(28)
        p1.font.bold = True
        p1.font.color.rgb = color
        p1.space_after = Pt(10)

        p2 = tf.add_paragraph()
        p2.text = title
        p2.font.name = FONT_HEADING
        p2.font.size = Pt(16)
        p2.font.bold = True
        p2.font.color.rgb = TEXT_WHITE
        p2.space_after = Pt(4)

        p3 = tf.add_paragraph()
        p3.text = sub
        p3.font.name = FONT_HEADING
        p3.font.size = Pt(11)
        p3.font.bold = True
        p3.font.color.rgb = color
        p3.space_after = Pt(14)

        p4 = tf.add_paragraph()
        p4.text = body
        p4.font.name = FONT_BODY
        p4.font.size = Pt(12)
        p4.font.color.rgb = TEXT_MUTED

    # ═════════════════════════════════════════════════════════════════════════
    # SLIDE 4: Architecture & Workflow Pipeline
    # ═════════════════════════════════════════════════════════════════════════
    s4 = prs.slides.add_slide(blank_layout)
    set_slide_background(s4)
    add_header(s4, "End-to-End Diagnostic & Remediation Pipeline", "SYSTEM ARCHITECTURE")

    steps = [
        ("1. INTAKE", "Source & Manifests", "Scans Python AST tree for imports.\nParses pyproject.toml, requirements.txt, setup.py, and Pipfile.", ACCENT_BLUE),
        ("2. FINGERPRINT", "Runtime Snapshot", "Captures target OS, Python version, virtualenv, site-packages, and native tools into env.json.", ACCENT_PURPLE),
        ("3. RESOLVE", "4-Layer Mapping", "Translates module import names to PyPI distributions (e.g. cv2 → opencv-python).", ACCENT_AMBER),
        ("4. DIAGNOSE", "Rule Engine Evaluation", "Executes Rules R1–R7.\nDetects undeclared, missing, conflicting, or incompatible requirements.", ACCENT_ROSE),
        ("5. REMEDIATE", "Safe Fix Execution", "Generates tiered action plan.\nManifest updates, pip installs, winget system tools, PATH fix.", ACCENT_GREEN),
    ]

    sw = Inches(2.2)
    gap = Inches(0.18)
    for i, (step_title, sub, desc, color) in enumerate(steps):
        sx = Inches(0.8 + i * (2.2 + 0.18))
        sy = Inches(1.8)
        add_card(s4, sx, sy, sw, Inches(3.2))

        # Accent header bar
        hbar = s4.shapes.add_shape(MSO_SHAPE.RECTANGLE, sx, sy, sw, Inches(0.35))
        hbar.fill.solid()
        hbar.fill.fore_color.rgb = color
        hbar.line.fill.background()

        ht = s4.shapes.add_textbox(sx, sy, sw, Inches(0.35))
        ht.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        hp = ht.text_frame.paragraphs[0]
        hp.alignment = PP_ALIGN.CENTER
        hp.text = step_title
        hp.font.name = FONT_HEADING
        hp.font.size = Pt(11)
        hp.font.bold = True
        hp.font.color.rgb = BG_DARK

        tb = s4.shapes.add_textbox(sx + Inches(0.15), sy + Inches(0.5), sw - Inches(0.3), Inches(2.5))
        tf = tb.text_frame
        tf.word_wrap = True

        p1 = tf.paragraphs[0]
        p1.text = sub
        p1.font.name = FONT_HEADING
        p1.font.size = Pt(13)
        p1.font.bold = True
        p1.font.color.rgb = TEXT_WHITE
        p1.space_after = Pt(8)

        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.name = FONT_BODY
        p2.font.size = Pt(11)
        p2.font.color.rgb = TEXT_MUTED

    # Bottom summary callout card
    b_card = add_card(s4, Inches(0.8), Inches(5.3), Inches(11.73), Inches(1.5))
    b_tb = s4.shapes.add_textbox(Inches(1.1), Inches(5.45), Inches(11.1), Inches(1.2))
    b_tf = b_tb.text_frame
    b_tf.word_wrap = True
    bp1 = b_tf.paragraphs[0]
    bp1.text = "Dual-Ended Verification (Local vs. Baseline / Reference Machine)"
    bp1.font.name = FONT_HEADING
    bp1.font.size = Pt(14)
    bp1.font.bold = True
    bp1.font.color.rgb = ACCENT_GREEN
    bp1.space_after = Pt(4)
    bp2 = b_tf.add_paragraph()
    bp2.text = "Parity does not just diagnose the machine you are standing on: you can run 'parity diagnose . --against env.json' to test your code against a remote production server, CI container, or teammate's recorded machine environment before pushing code."
    bp2.font.name = FONT_BODY
    bp2.font.size = Pt(12)
    bp2.font.color.rgb = TEXT_MUTED

    # ═════════════════════════════════════════════════════════════════════════
    # SLIDE 5: Diagnostic Rule Engine
    # ═════════════════════════════════════════════════════════════════════════
    s5 = prs.slides.add_slide(blank_layout)
    set_slide_background(s5)
    add_header(s5, "Diagnostic Rule Catalog: Precise Root Cause Detection", "RULE EVALUATOR")

    rules = [
        ("R1", "Undeclared Import", "HIGH / LOW", "An imported module is detected in Python AST code but missing from dependency manifests. Runtime imports produce HIGH confidence; optional/try-except imports produce LOW confidence.", "Adds package to requirements.txt or installs via pip.", ACCENT_ROSE),
        ("R2", "Missing / Mismatched Dependency", "HIGH", "A declared package in requirements.txt or pyproject.toml is either absent from site-packages or violates version constraints (e.g. requires >=2.0, but 1.8 installed).", "Automated pip install or package upgrade.", ACCENT_AMBER),
        ("R3", "Conflicting Pins", "MEDIUM", "Contradictory or disjoint constraints detected across manifests or requirements files (e.g. file A pins ==1.0 while file B pins ==2.0).", "Highlights exact line evidence for manual unification.", ACCENT_PURPLE),
        ("R4", "Python Incompatibility", "HIGH", "Project's 'requires-python' specifier (e.g. >=3.11) conflicts with the active or target Python interpreter version.", "Alerts developer to switch or provision compatible interpreter.", ACCENT_BLUE),
        ("R7", "Marker Evaluation", "LOW (INFO)", "Evaluates PEP 508 environment markers (sys_platform, python_version, platform_machine) against the target fingerprint.", "Validates platform-specific dependencies.", ACCENT_GREEN),
    ]

    rw = Inches(11.73)
    rh = Inches(0.92)
    for i, (rid, name, sev, desc, fix, color) in enumerate(rules):
        ry = Inches(1.6 + i * 1.05)
        add_card(s5, Inches(0.8), ry, rw, rh)

        # Rule ID Pill
        pill = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), ry + Inches(0.18), Inches(0.85), Inches(0.55))
        pill.fill.solid()
        pill.fill.fore_color.rgb = color
        pill.line.fill.background()
        ptf = pill.text_frame
        ptf.vertical_anchor = MSO_ANCHOR.MIDDLE
        pp = ptf.paragraphs[0]
        pp.alignment = PP_ALIGN.CENTER
        pp.text = rid
        pp.font.name = FONT_HEADING
        pp.font.size = Pt(13)
        pp.font.bold = True
        pp.font.color.rgb = BG_DARK

        # Details
        dtb = s5.shapes.add_textbox(Inches(2.0), ry + Inches(0.12), Inches(9.8), Inches(0.7))
        dtf = dtb.text_frame
        dtf.word_wrap = True
        dp1 = dtf.paragraphs[0]
        dp1.text = f"{name}   [{sev}]"
        dp1.font.name = FONT_HEADING
        dp1.font.size = Pt(13)
        dp1.font.bold = True
        dp1.font.color.rgb = TEXT_WHITE

        dp2 = dtf.add_paragraph()
        dp2.text = f"{desc} -> Fix: {fix}"
        dp2.font.name = FONT_BODY
        dp2.font.size = Pt(10.5)
        dp2.font.color.rgb = TEXT_MUTED

    # ═════════════════════════════════════════════════════════════════════════
    # SLIDE 6: 4-Layer Import-to-Package Resolver
    # ═════════════════════════════════════════════════════════════════════════
    s6 = prs.slides.add_slide(blank_layout)
    set_slide_background(s6)
    add_header(s6, "Solving the 'Import Name != Package Name' Problem", "RESOLVER ENGINE")

    layers = [
        ("LAYER 1", "Installed Environment", "Queries importlib.metadata.packages_distributions() from the active or target venv. Directly maps top-level modules to installed distribution names.", "High Confidence", ACCENT_GREEN),
        ("LAYER 2", "Curated Mapping Catalog", "Bundled mapping fixture for notoriously ambiguous packages:\ncv2 → opencv-python\nyaml → PyYAML\nPIL → Pillow\nsklearn → scikit-learn", "Medium Confidence", ACCENT_BLUE),
        ("LAYER 3", "PEP 503 Identity Heuristic", "Standard Python packaging normalization fallback: lowercasing, replacing underscores and dots with dashes (e.g. date_util → date-util).", "Low Confidence", ACCENT_AMBER),
        ("LAYER 4", "PyPI Online Wheel Metadata", "Safely queries files.pythonhosted.org (when --online is enabled). Verifies SHA-256 wheel hash, bounded zip read to prevent zip-bombs.", "High Confidence", ACCENT_PURPLE),
    ]

    for i, (lname, title, body, conf, color) in enumerate(layers):
        lx = Inches(0.8 + (i % 2) * 5.95)
        ly = Inches(1.7 + (i // 2) * 2.5)
        add_card(s6, lx, ly, Inches(5.7), Inches(2.2))

        # Badge
        b = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, lx + Inches(0.3), ly + Inches(0.25), Inches(1.1), Inches(0.32))
        b.fill.solid()
        b.fill.fore_color.rgb = color
        b.line.fill.background()
        btf = b.text_frame
        btf.vertical_anchor = MSO_ANCHOR.MIDDLE
        bp = btf.paragraphs[0]
        bp.alignment = PP_ALIGN.CENTER
        bp.text = lname
        bp.font.name = FONT_HEADING
        bp.font.size = Pt(10)
        bp.font.bold = True
        bp.font.color.rgb = BG_DARK

        # Text box
        tb = s6.shapes.add_textbox(lx + Inches(1.55), ly + Inches(0.2), Inches(3.9), Inches(1.8))
        tf = tb.text_frame
        tf.word_wrap = True

        p1 = tf.paragraphs[0]
        p1.text = title
        p1.font.name = FONT_HEADING
        p1.font.size = Pt(14)
        p1.font.bold = True
        p1.font.color.rgb = TEXT_WHITE
        p1.space_after = Pt(4)

        p2 = tf.add_paragraph()
        p2.text = body
        p2.font.name = FONT_BODY
        p2.font.size = Pt(11)
        p2.font.color.rgb = TEXT_MUTED
        p2.space_after = Pt(4)

        p3 = tf.add_paragraph()
        p3.text = f"Reliability: {conf}"
        p3.font.name = FONT_HEADING
        p3.font.size = Pt(10.5)
        p3.font.bold = True
        p3.font.color.rgb = color

    # ═════════════════════════════════════════════════════════════════════════
    # SLIDE 7: Deep Environment Fingerprinting & Native Tools
    # ═════════════════════════════════════════════════════════════════════════
    s7 = prs.slides.add_slide(blank_layout)
    set_slide_background(s7)
    add_header(s7, "Deterministic Fingerprinting: Beyond Python Packages", "ENVIRONMENT PROFILER")

    # Left card: Python & OS
    add_card(s7, Inches(0.8), Inches(1.7), Inches(5.7), Inches(5.1))
    ltb = s7.shapes.add_textbox(Inches(1.1), Inches(2.0), Inches(5.1), Inches(4.5))
    ltf = ltb.text_frame
    ltf.word_wrap = True

    lp1 = ltf.paragraphs[0]
    lp1.text = "Python Runtime & OS Introspection"
    lp1.font.name = FONT_HEADING
    lp1.font.size = Pt(18)
    lp1.font.bold = True
    lp1.font.color.rgb = ACCENT_BLUE
    lp1.space_after = Pt(14)

    items = [
        "OS & Kernel: Windows, Linux glibc/musl, macOS Darwin",
        "Architecture: x86_64, arm64, platform tags (PEP 425)",
        "Python Interpreter: Exact executable binary path & version",
        "Virtualenv Isolation: Detects pyvenv.cfg, active isolation status",
        "Installed Package Map: Canonical distribution name & exact version",
        "Redaction Engine: --redact flag sanitizes usernames and sensitive filepaths before team sharing",
    ]
    for item in items:
        p = ltf.add_paragraph()
        p.text = f"•  {item}"
        p.font.name = FONT_BODY
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_MUTED
        p.space_after = Pt(6)

    # Right card: Native System Tools
    add_card(s7, Inches(6.8), Inches(1.7), Inches(5.7), Inches(5.1))
    rtb = s7.shapes.add_textbox(Inches(7.1), Inches(2.0), Inches(5.1), Inches(4.5))
    rtf = rtb.text_frame
    rtf.word_wrap = True

    rp1 = rtf.paragraphs[0]
    rp1.text = "Native System Tools Detection"
    rp1.font.name = FONT_HEADING
    rp1.font.size = Pt(18)
    rp1.font.bold = True
    rp1.font.color.rgb = ACCENT_AMBER
    rp1.space_after = Pt(14)

    titems = [
        "Probed Tools: git, cmake, ffmpeg, docker, make, gcc, cl.exe (MSVC)",
        "PATH Verification: Differentiates FOUND, FOUND_NOT_ON_PATH, BROKEN, and NOT_FOUND",
        "CWE-426 Protection: Safe which probes avoid current working directory binary hijacking",
        "Execution Isolation: Tools probed in isolated temp directory with 5-second timeouts",
        "Automated Windows Winget: One-click background installation for missing system tools",
        "Dynamic PATH Registration: Windows registry + process PATH updates without reboot",
    ]
    for item in titems:
        p = rtf.add_paragraph()
        p.text = f"•  {item}"
        p.font.name = FONT_BODY
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_MUTED
        p.space_after = Pt(6)

    # ═════════════════════════════════════════════════════════════════════════
    # SLIDE 8: Interactive Web Dashboard
    # ═════════════════════════════════════════════════════════════════════════
    s8 = prs.slides.add_slide(blank_layout)
    set_slide_background(s8)
    add_header(s8, "Interactive Web Dashboard: Visual Health & 1-Click Fixes", "GRAPHICAL USER INTERFACE")

    features = [
        ("📊 Executive Dashboard", "Instant overview cards for Total Findings, Declared Dependencies, and AST Imported Modules with severity breakdowns.", ACCENT_BLUE),
        ("🔍 Findings & Code Evidence", "Expandable findings listing file paths, exact line numbers, explanation templates, and automated/manual fix proposals.", ACCENT_PURPLE),
        ("⚡ 1-Click Auto Solve", "Directly appends undeclared dependencies to requirements.txt or triggers safe background pip installations.", ACCENT_GREEN),
        ("🔧 Native Tools & Winget", "Real-time tool status badges with 1-click winget installer and dynamic Windows user PATH modification.", ACCENT_AMBER),
        ("🔗 Live Import Resolver", "Interactive test console with auto-suggest candidates, layer resolution paths, and confidence indicators.", ACCENT_ROSE),
        ("📄 Raw JSON & Snapshots", "Full transparency into the complete diagnostic payload and exportable environment fingerprint.", ACCENT_BLUE),
    ]

    for i, (title, desc, color) in enumerate(features):
        fx = Inches(0.8 + (i % 3) * 3.95)
        fy = Inches(1.7 + (i // 3) * 2.5)
        add_card(s8, fx, fy, Inches(3.8), Inches(2.2))

        # Accent dot
        dot = s8.shapes.add_shape(MSO_SHAPE.OVAL, fx + Inches(0.3), fy + Inches(0.3), Inches(0.18), Inches(0.18))
        dot.fill.solid()
        dot.fill.fore_color.rgb = color
        dot.line.fill.background()

        tb = s8.shapes.add_textbox(fx + Inches(0.6), fy + Inches(0.2), Inches(3.0), Inches(1.8))
        tf = tb.text_frame
        tf.word_wrap = True

        p1 = tf.paragraphs[0]
        p1.text = title
        p1.font.name = FONT_HEADING
        p1.font.size = Pt(14)
        p1.font.bold = True
        p1.font.color.rgb = TEXT_WHITE
        p1.space_after = Pt(8)

        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.name = FONT_BODY
        p2.font.size = Pt(11.5)
        p2.font.color.rgb = TEXT_MUTED

    # ═════════════════════════════════════════════════════════════════════════
    # SLIDE 9: CLI Developer Experience
    # ═════════════════════════════════════════════════════════════════════════
    s9 = prs.slides.add_slide(blank_layout)
    set_slide_background(s9)
    add_header(s9, "CLI-First Architecture: Powerful, Scriptable & Fast", "COMMAND-LINE INTERFACE")

    cli_cards = [
        ("1. Scan Project Dependencies", "parity scan ./my-project --all", "Inspects AST imports, parses manifests, and prints categorized runtime vs dev vs optional dependency matrices.", ACCENT_BLUE),
        ("2. Capture Machine Fingerprint", "parity fingerprint --out env.json --redact", "Captures deterministic JSON snapshot of active runtime, packages, and system compilers with sanitized paths.", ACCENT_PURPLE),
        ("3. Diagnose Against Target", "parity diagnose . --against env.json --format json", "Runs diagnostic engine against local system or recorded environment JSON. Supports machine-readable JSON for CI/CD.", ACCENT_GREEN),
        ("4. Resolve Import Name", "parity resolve cv2 --explain", "Tests multi-layer package mapping for any Python import with full layer explanation and confidence ratings.", ACCENT_AMBER),
    ]

    for i, (title, cmd, desc, color) in enumerate(cli_cards):
        cy = Inches(1.7 + i * 1.3)
        add_card(s9, Inches(0.8), cy, Inches(11.73), Inches(1.15))

        # Command terminal box inside card
        tb = s9.shapes.add_textbox(Inches(1.1), cy + Inches(0.12), Inches(11.1), Inches(0.95))
        tf = tb.text_frame
        tf.word_wrap = True

        p1 = tf.paragraphs[0]
        p1.text = f"{title}:   "
        p1.font.name = FONT_HEADING
        p1.font.size = Pt(13)
        p1.font.bold = True
        p1.font.color.rgb = TEXT_WHITE

        # Add inline code
        p2 = tf.add_paragraph()
        p2.text = f"$  {cmd}"
        p2.font.name = FONT_CODE
        p2.font.size = Pt(12)
        p2.font.bold = True
        p2.font.color.rgb = color
        p2.space_after = Pt(2)

        p3 = tf.add_paragraph()
        p3.text = desc
        p3.font.name = FONT_BODY
        p3.font.size = Pt(11)
        p3.font.color.rgb = TEXT_MUTED

    # ═════════════════════════════════════════════════════════════════════════
    # SLIDE 10: Security & Threat Model
    # ═════════════════════════════════════════════════════════════════════════
    s10 = prs.slides.add_slide(blank_layout)
    set_slide_background(s10)
    add_header(s10, "Built for Untrusted Codebases: Zero-Execution Threat Model", "SECURITY ARCHITECTURE")

    sec_cards = [
        ("Execution Isolation", "Parity never executes setup.py or untrusted target virtual environments. When querying sysconfig, Parity forces the Python binary to run with '-S -I', completely skipping site initialization and neutralizing .pth exploits.", ACCENT_GREEN),
        ("Binary Hijacking Defense (CWE-426)", "Safe tool probing uses sanitized PATH environments that explicitly block execution from the current project directory or cwd. Probe tests execute in ephemeral temporary directories with strict 5-second timeouts.", ACCENT_BLUE),
        ("Data-Only .pth Parsing", "When discovering editable installs from .pth files, Parity reads files as plain text, explicitly rejecting and ignoring lines containing 'import ' statements to prevent arbitrary script execution.", ACCENT_PURPLE),
        ("Online Wheel Integrity & DoS Defense", "Layer 4 online queries strictly use files.pythonhosted.org, compute SHA-256 hashes immediately upon streaming before opening zip archives, and enforce bounded reads to protect against zip-bombs.", ACCENT_ROSE),
    ]

    for i, (title, desc, color) in enumerate(sec_cards):
        sx = Inches(0.8 + (i % 2) * 5.95)
        sy = Inches(1.7 + (i // 2) * 2.5)
        add_card(s10, sx, sy, Inches(5.7), Inches(2.2))

        # Shield icon box
        sbox = s10.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, sx + Inches(0.3), sy + Inches(0.25), Inches(0.4), Inches(0.4))
        sbox.fill.solid()
        sbox.fill.fore_color.rgb = color
        sbox.line.fill.background()

        tb = s10.shapes.add_textbox(sx + Inches(0.9), sy + Inches(0.2), Inches(4.5), Inches(1.8))
        tf = tb.text_frame
        tf.word_wrap = True

        p1 = tf.paragraphs[0]
        p1.text = title
        p1.font.name = FONT_HEADING
        p1.font.size = Pt(14)
        p1.font.bold = True
        p1.font.color.rgb = TEXT_WHITE
        p1.space_after = Pt(6)

        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.name = FONT_BODY
        p2.font.size = Pt(11)
        p2.font.color.rgb = TEXT_MUTED

    # ═════════════════════════════════════════════════════════════════════════
    # SLIDE 11: Validation & Test Corpus
    # ═════════════════════════════════════════════════════════════════════════
    s11 = prs.slides.add_slide(blank_layout)
    set_slide_background(s11)
    add_header(s11, "Empirical Validation: Test Corpus & Determinism", "QUALITY & BENCHMARKS")

    metrics = [
        ("46", "Passing Test Cases", "Covering AST parsing, manifest extraction, environment fingerprinting, and resolution.", ACCENT_GREEN),
        ("100%", "Deterministic Runs", "Hypothesis property-based tests verify identical outputs regardless of run order.", ACCENT_BLUE),
        ("10+", "Test Corpus Scenarios", "Healthy projects, undeclared imports, version drift, missing tools, and broken markers.", ACCENT_PURPLE),
        ("< 1s", "Typical Scan Time", "Instantaneous AST parsing without spinning up slow package managers or interpreters.", ACCENT_AMBER),
    ]

    mw = Inches(2.78)
    for i, (stat, label, desc, color) in enumerate(metrics):
        mx = Inches(0.8 + i * 2.95)
        my = Inches(1.7)
        add_card(s11, mx, my, mw, Inches(2.3))

        tb = s11.shapes.add_textbox(mx + Inches(0.2), my + Inches(0.25), mw - Inches(0.4), Inches(1.8))
        tf = tb.text_frame
        tf.word_wrap = True

        p1 = tf.paragraphs[0]
        p1.text = stat
        p1.font.name = FONT_HEADING
        p1.font.size = Pt(36)
        p1.font.bold = True
        p1.font.color.rgb = color
        p1.space_after = Pt(4)

        p2 = tf.add_paragraph()
        p2.text = label
        p2.font.name = FONT_HEADING
        p2.font.size = Pt(12.5)
        p2.font.bold = True
        p2.font.color.rgb = TEXT_WHITE
        p2.space_after = Pt(4)

        p3 = tf.add_paragraph()
        p3.text = desc
        p3.font.name = FONT_BODY
        p3.font.size = Pt(10.5)
        p3.font.color.rgb = TEXT_MUTED

    # Bottom card: Test Suite Breakdown
    bc = add_card(s11, Inches(0.8), Inches(4.3), Inches(11.73), Inches(2.5))
    bt = s11.shapes.add_textbox(Inches(1.1), Inches(4.5), Inches(11.1), Inches(2.1))
    btf = bt.text_frame
    btf.word_wrap = True

    bp1 = btf.paragraphs[0]
    bp1.text = "Comprehensive Test Suite Architecture"
    bp1.font.name = FONT_HEADING
    bp1.font.size = Pt(16)
    bp1.font.bold = True
    bp1.font.color.rgb = TEXT_WHITE
    bp1.space_after = Pt(10)

    tests_list = [
        ("tests/test_determinism.py", "Verifies JSON fingerprint serialization determinism and key sorting stability."),
        ("tests/test_diagnose.py", "Runs corpus evaluation across healthy_*, broken_1 (undeclared), broken_2 (mismatch), and broken_8 (optional)."),
        ("tests/test_env.py", "Evaluates native tools detection, CWE-426 binary planting mitigation, and PATH sanitization."),
        ("tests/test_resolve.py", "Verifies 4-layer resolution logic, bundled mappings, PEP 503 normalization, and wheel validation."),
        ("tests/test_scan.py", "Tests AST import categorizer (Runtime, Dev, Optional) and manifest parsers (pyproject, requirements, setup.py)."),
    ]
    for name, detail in tests_list:
        p = btf.add_paragraph()
        p.text = f"•  {name}: {detail}"
        p.font.name = FONT_BODY
        p.font.size = Pt(11)
        p.font.color.rgb = TEXT_MUTED
        p.space_after = Pt(3)

    # ═════════════════════════════════════════════════════════════════════════
    # SLIDE 12: Roadmap & Conclusion
    # ═════════════════════════════════════════════════════════════════════════
    s12 = prs.slides.add_slide(blank_layout)
    set_slide_background(s12)
    add_header(s12, "Future Roadmap & Project Summary", "CONCLUSION & NEXT STEPS")

    # Left: Roadmap
    add_card(s12, Inches(0.8), Inches(1.7), Inches(5.7), Inches(5.1))
    ltb = s12.shapes.add_textbox(Inches(1.1), Inches(2.0), Inches(5.1), Inches(4.5))
    ltf = ltb.text_frame
    ltf.word_wrap = True

    lp1 = ltf.paragraphs[0]
    lp1.text = "Strategic Roadmap"
    lp1.font.name = FONT_HEADING
    lp1.font.size = Pt(18)
    lp1.font.bold = True
    lp1.font.color.rgb = ACCENT_PURPLE
    lp1.space_after = Pt(14)

    roadmap_items = [
        ("Automated Fix Sandbox", "Dry-run execution of pip installations inside temporary throwaway environments before committing to the main virtualenv."),
        ("IDE Language Server (LSP)", "Real-time editor diagnostics in VS Code: highlight undeclared imports instantly as you write code."),
        ("CI/CD GitHub Action", "Automated pull request gating: reject PRs containing undeclared imports or environment drift."),
        ("Multi-Ecosystem Expansion", "Extending Parity's AST and manifest engine to Node.js/npm and Rust/Cargo environments."),
    ]
    for r_title, r_desc in roadmap_items:
        p = ltf.add_paragraph()
        p.text = f"▸  {r_title}"
        p.font.name = FONT_HEADING
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = TEXT_WHITE
        p.space_after = Pt(2)

        p2 = ltf.add_paragraph()
        p2.text = r_desc
        p2.font.name = FONT_BODY
        p2.font.size = Pt(11)
        p2.font.color.rgb = TEXT_MUTED
        p2.space_after = Pt(8)

    # Right: Conclusion & Summary
    add_card(s12, Inches(6.8), Inches(1.7), Inches(5.7), Inches(5.1))
    rtb = s12.shapes.add_textbox(Inches(7.1), Inches(2.0), Inches(5.1), Inches(4.5))
    rtf = rtb.text_frame
    rtf.word_wrap = True

    rp1 = rtf.paragraphs[0]
    rp1.text = "Why Parity Matters"
    rp1.font.name = FONT_HEADING
    rp1.font.size = Pt(18)
    rp1.font.bold = True
    rp1.font.color.rgb = ACCENT_GREEN
    rp1.space_after = Pt(14)

    summary_points = [
        "Eliminates 'Works on My Machine': Replaces hours of frantic environment debugging with deterministic, evidence-backed explanations.",
        "Zero Security Compromise: Never executes untrusted repository code to inspect it, neutralizing .pth and binary planting attack vectors.",
        "Developer-Friendly: Provides both a scriptable, JSON-ready CLI for power users and CI/CD pipelines, and a modern GUI for one-click fixes.",
        "Production Ready: Built with strict typing (mypy), Ruff linting, Hatchling packaging, and 46 passing tests.",
    ]
    for sp in summary_points:
        p = rtf.add_paragraph()
        p.text = f"✓  {sp}"
        p.font.name = FONT_BODY
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_MUTED
        p.space_after = Pt(10)

    return prs

def main():
    prs = create_presentation()
    out_path = Path("Parity_Presentation.pptx").resolve()
    prs.save(str(out_path))
    print(f"Presentation saved successfully to: {out_path}")

if __name__ == "__main__":
    main()
